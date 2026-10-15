import sys, csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
S, out = sys.argv[1], sys.argv[2]
notes = {
 "DS-Veranstaltungstechnik GmbH": "E-Mail-Domain ds-veranstaltung.de",
 "PA-Technik-Team Veranstaltungstechnik": "Nur Postfach bekannt – Adresse telefonisch klären (07063 2679658)",
 "EuroTronic Veranstaltungstechnik UG": "Verzeichnisse nennen Schwimmbadweg 46 – vor Versand prüfen",
 "Chris Only Entertainment / TONLYGHTS Event Technik": "Website nennt auch Giebelstadt und Wittighausen",
 "Maximilian Egner Veranstaltungstechnik": "Website im Umbau",
 "Kuehner Events GbR": "Tritt als Betrieb aus Kirchheim/Teck auf",
 "AG Veranstaltungstechnik": "Andere Seiten nennen Albstraße 1, 73274 Notzingen; evtl. selber Inhaber wie UBS Quality Sound",
 "UBS Quality Sound (Alexander Fischer)": "Evtl. selber Inhaber wie AG Veranstaltungstechnik",
 "Lite-Tech Veranstaltungstechnik GmbH": "Adresse aus Mietkatalog 2023",
 "Soundlight Company GmbH": "Büro Planckstr. 21",
 "wivent (Andrea Wilhelm)": "Adresse von Startseite, Impressum nicht gelesen",
 "HZ Eventtechnik": "Adresse von Startseite, Impressum nicht gelesen",
 "Meier Eventservice": "Impressum-Stand ca. 4,5 Jahre alt – vor Versand prüfen",
 "BW Sound & Light (Benjamin Weschenfelder)": "Ältere Seite nennt Schumannstraße 30",
 "da-sound PA-Verleih & Eventservice": "Vermietet Benzin-Aggregate (Top-Lead); Lager Werner-von-Siemens-Str. 2, Pfungstadt",
 "LSD Ton- & Lichtverleih": "Impressum nennt auch Kirchstr. 29, 68623 Lampertheim",
 "Soundstation Freiburg": "Büro; Lager/Abholung Wiesentalstr. 29, 79115 Freiburg",
 "H-Light Veranstaltungstechnik (Oliver Hub)": "Adresse von Kontaktseite (Lager); Büro Am Wall 36, 97525 Schwebheim",
 "ICEVT Event & Veranstaltungstechnik": "Kein Impressum gefunden – Adresse aus Websitedaten",
 "DAXL Veranstaltungstechnik e.K.": "Lager Mornewegstr. 43B",
 "Event-Xpress.com GmbH": "Büro; Sitz laut Impressum Hindenburgstr. 6a",
 "Beattec Veranstaltungstechnik (Jonas Baur)": "LK Lindau (Allgäu), Randgebiet",
 "AE Audio-Express Veranstaltungstechnik GmbH & Co. KG": "Impressum-Stand 1-4 Jahre alt – vor Versand prüfen",
 "Prolite Event GmbH": "Impressum-Stand 1-4 Jahre alt – vor Versand prüfen",
 "Hübsch Licht Ton Events (Peter Hübsch)": "Impressum-Stand 1-4 Jahre alt – vor Versand prüfen",
 "BMS Veranstaltungstechnik GmbH": "Impressum-Stand 1-4 Jahre alt – vor Versand prüfen",
 "C2 CONCERTS GmbH": "Datenschutzseite nennt Nr. 28 – Impressum Nr. 8",
 "Göckelesmaier Festbetriebs GmbH": "Andere Seite nennt Nr. 76",
 "KARO Konzert-Agentur Rothenburg GmbH": "Ältere Seite nennt Schweinsdorfer Str. 25b",
 "MotorCore GmbH (Foodtruckfabrik)": "Kontaktseite nennt 19H",
 "Losberger GmbH": "Adresse aus Produktdatenblatt, nicht Impressum",
 "Procon Event GmbH (Procon Veranstaltungstechnik)": "Verzeichnisse nennen noch Frangenheimstraße",
 "satis&fy AG": "Keine Hausnummer im Impressum",
 "Presented for People GmbH & Co. KG (Rocco del Schlacko)": "Festival laut Website beendet – schwacher Lead",
 "Reload Event GmbH": "Adresse aus Shop-AGB/Datenschutz",
 "Otter Veranstaltungstechnik GbR": "Laut Website keine Abholadresse",
 "ZELTCITY GmbH & Co. KG": "PLZ nicht direkt im Impressum bestätigt",
 "Kopf & Steine GmbH (MS Dockville)": "Impressum evtl. älter",
 "GATE Eventmanagement & Veranstaltungstechnik GmbH": "Kontaktseite nennt auch Reuchlinstr. 10-11, 10553 Berlin",
 "DTB Veranstaltungstechnik GmbH": "Lager: Am Buchhorst 33, 14478 Potsdam",
 "ESS Emmes Sound Service": "Adresse am schwächsten belegt – vor Versand prüfen",
 "Expovent": "Impressum ca. 6 Jahre alt – vor Versand prüfen",
}
import glob, os
from urllib.parse import urlparse
REG = {"Baden-Württemberg":"Süd","Bayern":"Süd","Nordrhein-Westfalen":"West","Hessen":"West","Rheinland-Pfalz":"West","Saarland":"West",
       "Hamburg":"Nord","Bremen":"Nord","Niedersachsen":"Nord","Schleswig-Holstein":"Nord","Mecklenburg-Vorpommern":"Nord",
       "Berlin":"Ost","Brandenburg":"Ost","Sachsen":"Ost","Sachsen-Anhalt":"Ost","Thüringen":"Ost"}
ORDER = {"Süd":0,"West":1,"Nord":2,"Ost":3}
EXCL = ("lightequip","rentboks","lautmacher","huss-licht","audiopro")
files = [f"{S}/{f}.txt" for f in ("sued","lokal1","lokal2","west","nord","ost")] + sorted(glob.glob(f"{S}/neu/*.txt"))
wb = Workbook(); ws = wb.active; ws.title = "Adressliste"
hdr = ["Nr.","Firma","Straße","PLZ","Ort","Bundesland","Region","Kategorie","Website","Warum Bedarf (PowerRock / Delta Pro Ultra)","Hinweis zur Adresse","Priorität"]
ws.append(hdr)
seen=set(); rows=[]; skipped=[]
def key(url,name):
    h=urlparse(url if "//" in url else "http://"+url).netloc.lower()
    return h[4:] if h.startswith("www.") else h or name.lower()
for f in files:
    for line in open(f, encoding="utf-8"):
        line=line.strip()
        if not line or line.startswith("#"): continue
        p=[x.strip() for x in line.split(";")]
        if len(p) not in (8,9): skipped.append((os.path.basename(f),line[:60])); continue
        if len(p)==8: p.append(notes.get(p[0],""))
        p[2]=p[2].zfill(5)
        k=key(p[6],p[0])
        if k in seen or any(e in k for e in EXCL) or p[4] not in REG: skipped.append((os.path.basename(f),p[0])); continue
        import re as _re
        k2=p[2]+_re.sub(r"[^a-z0-9]","",p[1].lower().replace("straße","str").replace("strasse","str").replace("str.","str"))
        if k2 in seen: skipped.append((os.path.basename(f),"dup-addr "+p[0])); continue
        seen.add(k); seen.add(k2)
        rows.append(p)
import re as _re2
def prio(p):
    t=" ".join(p[:9]).lower()
    if _re2.search(r"konferenz|kongress|congress|installation|mietmobiliar|nur bedingt|fit fraglich|technikanteil unklar|verleih .*nicht belegt|nicht bestätigt|weder bestätigt", t): return "C"
    if _re2.search(r"aggregat|generator|festival|open.air|stadtfest|zelt|bühne|buehne|foodtruck|streetfood|street.food|eventstrom|mobile bar|kaffeemobil|outdoor|volksfest|weinfest|weihnachtsmarkt", t): return "A"
    return "B"
rows.sort(key=lambda p:(ORDER[REG[p[4]]],p[4],p[2]))
n=0
for p in rows:
    n+=1
    ws.append([n,p[0],p[1],p[2],p[3],p[4],REG[p[4]],p[5],p[6],p[7],p[8],prio(p)])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
    c.alignment = Alignment(vertical="center", wrap_text=True)
for row in ws.iter_rows(min_row=2):
    row[3].number_format = "@"
    row[8].hyperlink = row[8].value; row[8].font = Font(color="0563C1", underline="single")
widths = [5,42,30,7,24,20,7,30,38,60,45,9]
for i, w in enumerate(widths): ws.column_dimensions[chr(65+i)].width = w
ws.freeze_panes = "C2"
t = Table(displayName="Adressen", ref=f"A1:L{n+1}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
ws.add_table(t)
lg = wb.create_sheet("Legende")
for row in [["Priorität","Bedeutung"],
            ["A","Klarer Outdoor-/Netzersatz-Bedarf: Festival-/Open-Air-/Stadtfest-Veranstalter, Zelt- und Bühnenverleih, Foodtrucks, mobile Bars, Firmen mit Aggregaten"],
            ["B","Passende Zielgruppe ohne konkreten Outdoor-Hinweis: Veranstaltungstechnik, Eventagenturen, Event-Catering"],
            ["C","Schwächere Passung oder ungeklärt: Konferenz-/Kongresstechnik, Installation, Mietmobiliar, Firmen die sich per Recherche nicht bestätigen ließen"],
            [],["Hinweis","Quelle der Adresse (eigene Website / Branchenverzeichnis) und Prüfvermerke"]]:
    lg.append(row)
lg.column_dimensions["A"].width=12; lg.column_dimensions["B"].width=120
lg["A1"].font=Font(bold=True); lg["B1"].font=Font(bold=True)
wb.save(out)
# CSV for mail merge
with open(out.replace(".xlsx",".csv"), "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh, delimiter=";")
    for r in ws.iter_rows(values_only=True): w.writerow(r)
print(n); print("skipped:",skipped)
